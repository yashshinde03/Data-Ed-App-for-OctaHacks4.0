#!/usr/bin/env python3
"""
generate_docx.py
----------------
Generates docs/assignment.docx from the assignment content.

Requirements:
    pip install python-docx pillow

Usage:
    python3 docs/generate_docx.py

The generated file is written to docs/assignment.docx in the repository root.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(SCRIPT_DIR)
ER_IMAGE = os.path.join(SCRIPT_DIR, "er_diagram_placeholder.png")
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "assignment.docx")

# ---------------------------------------------------------------------------
# Helper utilities
# ---------------------------------------------------------------------------

def add_horizontal_rule(doc: Document) -> None:
    """Add a thin horizontal line paragraph."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_code_block(doc: Document, code: str) -> None:
    """Add a monospaced code block paragraph with light-grey shading."""
    for line in code.strip().split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        # Light-grey background shading on the paragraph
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)

        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)

    # Blank gap after code block
    doc.add_paragraph()


def set_heading_color(para, r: int, g: int, b: int) -> None:
    for run in para.runs:
        run.font.color.rgb = RGBColor(r, g, b)


# ---------------------------------------------------------------------------
# SQL content
# ---------------------------------------------------------------------------

SQL_CREATE = """\
CREATE DATABASE fifo_db;
USE fifo_db;

CREATE TABLE Employee (
  employee_id INT AUTO_INCREMENT,
  name        VARCHAR(100) NOT NULL,
  role        VARCHAR(50),
  contact     VARCHAR(100),
  PRIMARY KEY (employee_id)
);

CREATE TABLE Site (
  site_id   INT AUTO_INCREMENT,
  site_name VARCHAR(100) NOT NULL,
  location  VARCHAR(100),
  PRIMARY KEY (site_id)
);

CREATE TABLE Room (
  room_id  INT AUTO_INCREMENT,
  capacity INT CHECK (capacity > 0),
  site_id  INT,
  PRIMARY KEY (room_id),
  FOREIGN KEY (site_id) REFERENCES Site(site_id)
);

CREATE TABLE Shift (
  shift_id   INT AUTO_INCREMENT,
  start_date DATE NOT NULL,
  end_date   DATE NOT NULL,
  PRIMARY KEY (shift_id)
);

CREATE TABLE Allocation (
  allocation_id INT AUTO_INCREMENT,
  employee_id   INT,
  room_id       INT,
  shift_id      INT,
  PRIMARY KEY (allocation_id),
  FOREIGN KEY (employee_id) REFERENCES Employee(employee_id),
  FOREIGN KEY (room_id)     REFERENCES Room(room_id),
  FOREIGN KEY (shift_id)    REFERENCES Shift(shift_id)
);"""

SQL_INSERT = """\
INSERT INTO Employee (name, role, contact)
VALUES ('Alex Johnson', 'Engineer',   '123456'),
       ('Sarah Lee',    'Technician', '789101');

INSERT INTO Site (site_name, location)
VALUES ('Site A', 'WA'),
       ('Site B', 'WA');

INSERT INTO Room (capacity, site_id)
VALUES (2, 1),
       (1, 2);

INSERT INTO Shift (start_date, end_date)
VALUES ('2026-04-01', '2026-04-14'),
       ('2026-04-15', '2026-04-28');

INSERT INTO Allocation (employee_id, room_id, shift_id)
VALUES (1, 1, 1),
       (2, 2, 2);"""

SQL_QUERIES = """\
-- Retrieve all employees
SELECT * FROM Employee;

-- Join: employee name, site and shift start date
SELECT e.name,
       s.site_name,
       sh.start_date
FROM   Allocation a
JOIN   Employee e  ON a.employee_id = e.employee_id
JOIN   Room     r  ON a.room_id     = r.room_id
JOIN   Site     s  ON r.site_id     = s.site_id
JOIN   Shift    sh ON a.shift_id    = sh.shift_id;

-- Aggregation: total rooms per site
SELECT site_id,
       COUNT(room_id) AS total_rooms
FROM   Room
GROUP  BY site_id;

-- Aggregation: total employees per role
SELECT role,
       COUNT(*) AS total_employees
FROM   Employee
GROUP  BY role;"""


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------

def build_document() -> None:
    doc = Document()

    # -----------------------------------------------------------------------
    # Page margins
    # -----------------------------------------------------------------------
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    # -----------------------------------------------------------------------
    # Title block
    # -----------------------------------------------------------------------
    title = doc.add_heading("INMT5526: Business Intelligence", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Individual Assignment – Semester 1, 2026")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].bold = True

    meta = doc.add_paragraph(
        "University of Western Australia Business School – Management and Organisations\n"
        "CRICOS: 00126G | PRV12169, Australian University"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_horizontal_rule(doc)

    # -----------------------------------------------------------------------
    # 1. Introduction
    # -----------------------------------------------------------------------
    h1 = doc.add_heading("1. Introduction", level=1)
    set_heading_color(h1, 31, 73, 125)

    doc.add_paragraph(
        "This report focuses on designing and implementing a database solution to address a "
        "business problem related to workforce and accommodation management in FIFO "
        "(Fly-In Fly-Out) operations. FIFO systems are commonly used in industries such as "
        "mining in Western Australia, where workers are transported to remote locations for "
        "specific time periods."
    )
    doc.add_paragraph(
        "The main objective of this assignment is to apply database concepts learned in the "
        "first four topics of the unit, including entity-relationship modelling and SQL, to "
        "develop a structured solution. The report outlines the problem, explains the design "
        "process, presents the implementation and reflects on the effectiveness of the solution."
    )

    add_horizontal_rule(doc)

    # -----------------------------------------------------------------------
    # 2. Background
    # -----------------------------------------------------------------------
    h2 = doc.add_heading("2. Background", level=1)
    set_heading_color(h2, 31, 73, 125)

    doc.add_paragraph(
        "FIFO operations involve managing multiple components at the same time, including "
        "employees, accommodation, shifts and site locations. In many organisations, this "
        "process is still handled using spreadsheets or basic systems, which can create "
        "several issues such as data duplication, errors and inefficiencies."
    )
    doc.add_paragraph(
        "A key problem is the lack of a centralised and structured system to manage workforce "
        "allocation and accommodation. For example, employees may be assigned to overlapping "
        "shifts or rooms may be overbooked due to poor tracking. These issues can affect "
        "operational efficiency and increase costs."
    )
    doc.add_paragraph(
        "Databases provide a more reliable solution because they allow data to be stored in a "
        "structured format and accessed efficiently. As discussed in the unit, databases are "
        "designed to organise data in a consistent way so that it can be retrieved and "
        "analysed to support decision-making (Ramakrishnan & Gehrke, 2003). Compared to "
        "spreadsheets, databases are more suitable when dealing with large amounts of data "
        "and multiple users."
    )
    doc.add_paragraph(
        "The proposed solution is to develop a relational database system that stores "
        "information about employees, shifts, rooms and site locations. This system would "
        "allow managers to easily track which employees are assigned to which shifts and "
        "where they are staying."
    )
    doc.add_paragraph(
        "This problem is worth solving because it improves coordination, reduces errors and "
        "supports better decision-making. From a business intelligence perspective, having "
        "structured and reliable data is essential for generating insights and making "
        "informed decisions (Negash, 2004)."
    )

    add_horizontal_rule(doc)

    # -----------------------------------------------------------------------
    # 3. Design
    # -----------------------------------------------------------------------
    h3 = doc.add_heading("3. Design", level=1)
    set_heading_color(h3, 31, 73, 125)

    doc.add_paragraph(
        "The database design was developed using entity-relationship modelling, which is used "
        "to represent real-world objects as entities and define the relationships between them. "
        "An entity refers to something we store data about, such as an employee or a site, "
        "while attributes describe the properties of that entity (Chen, 1976)."
    )

    # Entities and Attributes
    h3a = doc.add_heading("Entities and Attributes", level=2)
    set_heading_color(h3a, 68, 114, 196)

    doc.add_paragraph("The following entities were identified for the system:")

    entities = {
        "Employee": ["employee_id (Primary Key)", "name", "role", "contact"],
        "Site": ["site_id (Primary Key)", "site_name", "location"],
        "Room": ["room_id (Primary Key)", "capacity", "site_id (Foreign Key → Site)"],
        "Shift": ["shift_id (Primary Key)", "start_date", "end_date"],
        "Allocation": [
            "allocation_id (Primary Key)",
            "employee_id (Foreign Key → Employee)",
            "room_id (Foreign Key → Room)",
            "shift_id (Foreign Key → Shift)",
        ],
    }

    for entity, attrs in entities.items():
        p = doc.add_paragraph()
        run = p.add_run(entity)
        run.bold = True
        run.underline = True
        for attr in attrs:
            bullet = doc.add_paragraph(attr, style="List Bullet")
            bullet.paragraph_format.left_indent = Inches(0.5)

    # Relationships
    h3b = doc.add_heading("Relationships", level=2)
    set_heading_color(h3b, 68, 114, 196)

    doc.add_paragraph(
        "The relationships between entities were defined as follows:"
    )
    relationships = [
        "A site can contain multiple rooms (one-to-many)",
        "An employee can have multiple allocations (one-to-many)",
        "A room can be assigned to multiple allocations over time",
        "A shift can include multiple employees through allocations",
    ]
    for rel in relationships:
        doc.add_paragraph(rel, style="List Bullet")

    doc.add_paragraph(
        "These relationships reflect how FIFO operations function in reality and allow the "
        "system to store data efficiently."
    )

    # Design Justification
    h3c = doc.add_heading("Design Justification", level=2)
    set_heading_color(h3c, 68, 114, 196)

    doc.add_paragraph(
        "The design follows normalisation principles, meaning that data is organised into "
        "separate tables to reduce redundancy. Each table represents a single entity and "
        "relationships are managed through foreign keys."
    )
    doc.add_paragraph(
        "This approach ensures data consistency and avoids duplication. For example, employee "
        "details are stored only once and referenced where needed, rather than being repeated "
        "multiple times."
    )
    doc.add_paragraph(
        "Another important aspect of the design is planning before implementation. As "
        "highlighted in the lectures, designing the structure first helps avoid issues when "
        "writing SQL queries later."
    )

    # Use of Unit Concepts
    h3d = doc.add_heading("Use of Unit Concepts", level=2)
    set_heading_color(h3d, 68, 114, 196)

    doc.add_paragraph("The design applies several concepts from the unit:")
    for concept in [
        "Entity-relationship modelling",
        "Primary and foreign keys",
        "Structured relational data",
        "Planning before implementation",
    ]:
        doc.add_paragraph(concept, style="List Bullet")

    doc.add_paragraph(
        "These concepts were essential in ensuring that the database accurately represents "
        "the business problem. The ER diagram is provided in Appendix A."
    )

    add_horizontal_rule(doc)

    # -----------------------------------------------------------------------
    # 4. Solution
    # -----------------------------------------------------------------------
    h4 = doc.add_heading("4. Solution", level=1)
    set_heading_color(h4, 31, 73, 125)

    doc.add_paragraph(
        "The solution was implemented using SQL, which is used to create, manage and query "
        "relational databases. The implementation follows CRUD operations (Create, Read, "
        "Update and Delete), which form the basis of database systems (Elmasri & Navathe, 2016)."
    )

    h4a = doc.add_heading("Database and Table Creation", level=2)
    set_heading_color(h4a, 68, 114, 196)

    doc.add_paragraph(
        "The SQL tables were created to match the entities and relationships defined in the "
        "ER diagram (see Appendix A). Foreign key constraints enforce referential integrity "
        "between tables, ensuring that no allocation can reference a non-existent employee, "
        "room or shift."
    )

    h4b = doc.add_heading("Data Insertion", level=2)
    set_heading_color(h4b, 68, 114, 196)

    doc.add_paragraph(
        "Sample data was inserted to demonstrate that the system functions correctly. "
        "Two employees, two sites, two rooms, two shifts and two allocations were added "
        "as initial test records."
    )

    h4c = doc.add_heading("Queries", level=2)
    set_heading_color(h4c, 68, 114, 196)

    doc.add_paragraph(
        "Several queries were written to retrieve meaningful data from the database:"
    )
    for item in [
        "A simple SELECT query retrieves all employee records.",
        "A join query links Allocation, Employee, Room, Site and Shift tables to show "
        "each employee's assigned site and shift start date.",
        "An aggregation query counts the number of rooms per site using GROUP BY.",
        "An additional aggregation counts employees by role.",
    ]:
        doc.add_paragraph(item, style="List Number")

    doc.add_paragraph(
        "These queries demonstrate how structured data supports reporting and business "
        "intelligence. The full SQL statements are provided in Appendix B."
    )

    h4d = doc.add_heading("Explanation", level=2)
    set_heading_color(h4d, 68, 114, 196)

    doc.add_paragraph(
        "The SQL implementation demonstrates how data can be created, stored and retrieved "
        "efficiently. The use of joins allows multiple tables to be combined, while "
        "aggregation functions help generate useful insights."
    )
    doc.add_paragraph(
        "This reflects how databases support business intelligence by transforming raw data "
        "into meaningful information."
    )

    add_horizontal_rule(doc)

    # -----------------------------------------------------------------------
    # 5. Conclusion
    # -----------------------------------------------------------------------
    h5 = doc.add_heading("5. Conclusion", level=1)
    set_heading_color(h5, 31, 73, 125)

    doc.add_paragraph(
        "Overall, the database solution successfully addresses the problem of managing FIFO "
        "workforce and accommodation. By organising data into structured tables and defining "
        "clear relationships, the system improves efficiency and reduces the likelihood of errors."
    )
    doc.add_paragraph(
        "The solution meets the requirements by ensuring data consistency, enabling efficient "
        "queries and supporting decision-making. However, the system is still relatively basic "
        "and could be improved by adding features such as automated scheduling or integration "
        "with business intelligence tools."
    )
    doc.add_paragraph(
        "One challenge during the process was ensuring that relationships between tables were "
        "correctly defined, particularly when working with foreign keys. This was resolved by "
        "carefully reviewing the ER design before implementation."
    )
    doc.add_paragraph(
        "If this project were to be completed again, I would focus on expanding the system "
        "further and including more advanced queries to generate deeper insights."
    )

    add_horizontal_rule(doc)

    # -----------------------------------------------------------------------
    # Appendix A – ER Diagram
    # -----------------------------------------------------------------------
    doc.add_page_break()
    ha = doc.add_heading("Appendix A: Design (ER Diagram)", level=1)
    set_heading_color(ha, 31, 73, 125)

    notice_para = doc.add_paragraph()
    notice_run = notice_para.add_run(
        "NOTE: The image below is a placeholder. Replace it with the final ER Diagram "
        "exported from draw.io, Microsoft PowerPoint, or a similar diagramming tool "
        "before final submission."
    )
    notice_run.bold = True
    notice_run.font.color.rgb = RGBColor(192, 0, 0)

    if os.path.exists(ER_IMAGE):
        doc.add_picture(ER_IMAGE, width=Inches(6))
        caption = doc.add_paragraph(
            "Figure 1. Entity-Relationship Diagram for the FIFO Workforce Management Database."
        )
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            run.italic = True
    else:
        doc.add_paragraph(
            "[ER Diagram image not found. Please add er_diagram_placeholder.png to the "
            "docs/ directory and re-run generate_docx.py]"
        )

    # -----------------------------------------------------------------------
    # Appendix B – SQL Statements
    # -----------------------------------------------------------------------
    doc.add_page_break()
    hb = doc.add_heading("Appendix B: Implementation (SQL Statements)", level=1)
    set_heading_color(hb, 31, 73, 125)

    doc.add_paragraph(
        "All SQL statements below are formatted in Courier New (monospaced) font. "
        "They should be executed in the order presented: table creation, data insertion, "
        "then queries."
    )

    hb1 = doc.add_heading("Database and Table Creation", level=2)
    set_heading_color(hb1, 68, 114, 196)
    add_code_block(doc, SQL_CREATE)

    hb2 = doc.add_heading("Data Insertion", level=2)
    set_heading_color(hb2, 68, 114, 196)
    add_code_block(doc, SQL_INSERT)

    hb3 = doc.add_heading("Queries", level=2)
    set_heading_color(hb3, 68, 114, 196)
    add_code_block(doc, SQL_QUERIES)

    # -----------------------------------------------------------------------
    # References (APA)
    # -----------------------------------------------------------------------
    doc.add_page_break()
    hr = doc.add_heading("References", level=1)
    set_heading_color(hr, 31, 73, 125)

    references = [
        (
            "Chen, P. P. (1976). The entity-relationship model—toward a unified view of "
            "data. ",
            "ACM Transactions on Database Systems, 1",
            "(1), 9–36. https://doi.org/10.1145/320434.320440",
        ),
        (
            "Elmasri, R., & Navathe, S. B. (2016). ",
            "Fundamentals of database systems",
            " (7th ed.). Pearson.",
        ),
        (
            "Negash, S. (2004). Business intelligence. ",
            "Communications of the Association for Information Systems, 13",
            "(1), 177–195. https://doi.org/10.17705/1CAIS.01315",
        ),
        (
            "Ramakrishnan, R., & Gehrke, J. (2003). ",
            "Database management systems",
            " (3rd ed.). McGraw-Hill.",
        ),
    ]

    for normal_start, italic_part, normal_end in references:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)  # hanging indent
        p.paragraph_format.space_after = Pt(6)
        r1 = p.add_run(normal_start)
        r2 = p.add_run(italic_part)
        r2.italic = True
        r3 = p.add_run(normal_end)
        _ = r1, r3  # used implicitly via paragraph

    # -----------------------------------------------------------------------
    # Save
    # -----------------------------------------------------------------------
    doc.save(OUTPUT_PATH)
    print(f"Document saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_document()
